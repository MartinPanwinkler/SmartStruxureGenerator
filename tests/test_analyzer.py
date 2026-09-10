from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Border, Side

from template_analyzer import analyze_template


def test_detects_visual_blocks_and_merges(tmp_path: Path) -> None:
    path = tmp_path / "master.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Boxes"
    thin = Side(style="thin")
    for row in (1, 2, 5, 6):
        for col in (1, 2):
            sheet.cell(row, col).border = Border(left=thin, right=thin, top=thin, bottom=thin)
    sheet.merge_cells("A1:B1")
    sheet["A1"] = "Header"
    workbook.save(path)
    report = analyze_template(path)["sheets"]["Boxes"]
    assert report["candidate_boxes"] == ["A1:B2", "A5:B6"]
    assert report["merged_cells"] == ["A1:B1"]
