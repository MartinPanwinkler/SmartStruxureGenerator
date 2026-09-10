from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Font, PatternFill, Side

from generator import generate_workbook
from models import Record


def make_master(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Master"
    sheet.merge_cells("A1:D1")
    sheet["A1"] = "{{ASP}} {{MODULE_TYPE}}"
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A1"].fill = PatternFill("solid", fgColor="FFFF00")
    thin = Side(style="thin", color="000000")
    for row in range(1, 5):
        for col in range(1, 5):
            sheet.cell(row, col).border = Border(left=thin, right=thin, top=thin, bottom=thin)
    sheet.row_dimensions[1].height = 25
    sheet.column_dimensions["B"].width = 22
    workbook.save(path)


def configuration() -> dict:
    return {
        "group_by": ["asp", "module", "module_type"],
        "box_row_spacing": 2,
        "default_template": "standard",
        "templates": {"standard": {
            "sheet": "Master", "range": "A1:D4", "header_cells": {},
            "data_start_row": 2, "max_channels": 3,
            "columns": {"channel": "A", "datapoint": "B", "description": "C", "source": "D"},
        }},
    }


def test_copies_complete_box_and_splits(tmp_path: Path) -> None:
    master = tmp_path / "master.xlsx"
    output = tmp_path / "output.xlsx"
    make_master(master)
    records = [Record(asp="ASP-1", module="M1", module_type="DI", channel=str(i), datapoint=f"DP{i}") for i in range(5)]
    result = generate_workbook(master, output, records, configuration())
    assert result.boxes_created == 2
    book = load_workbook(output)
    sheet = book.active
    assert str(sheet.merged_cells) == "A1:D1 A7:D7"
    assert sheet["A1"].value == "ASP-1 DI"
    assert sheet["B2"].value == "DP0"
    assert sheet["B8"].value == "DP3"
    assert sheet["A1"].font.bold
    assert sheet["A1"].fill.fgColor.rgb == "00FFFF00"
    assert sheet.row_dimensions[1].height == 25
    assert sheet.column_dimensions["B"].width == 22
    book.close()


def test_never_overwrites_master(tmp_path: Path) -> None:
    master = tmp_path / "master.xlsx"
    make_master(master)
    before = master.read_bytes()
    with pytest.raises(ValueError, match="nicht überschrieben"):
        generate_workbook(master, master, [], configuration())
    assert master.read_bytes() == before


def test_refuses_existing_output(tmp_path: Path) -> None:
    master = tmp_path / "master.xlsx"
    output = tmp_path / "exists.xlsx"
    make_master(master)
    output.write_bytes(b"keep")
    with pytest.raises(FileExistsError):
        generate_workbook(master, output, [], configuration())
    assert output.read_bytes() == b"keep"


def test_explicit_two_column_slots(tmp_path: Path) -> None:
    master = tmp_path / "master.xlsx"
    output = tmp_path / "slots.xlsx"
    make_master(master)
    config = configuration()
    template = config["templates"]["standard"]
    template.pop("columns")
    template["max_channels"] = 2
    template["data_slots"] = [
        {"channel": "A2", "description": "B2"},
        {"channel": "C2", "description": "D2"},
    ]
    template["header_cells"] = {"module_header": ["A1"]}
    template["header_formats"] = {"module_header": "{module} {module_type}"}
    records = [
        Record(module="M1", module_type="DI", channel="01", description="Links"),
        Record(module="M1", module_type="DI", channel="02", description="Rechts"),
    ]
    generate_workbook(master, output, records, config)
    book = load_workbook(output)
    sheet = book.active
    assert sheet["A1"].value == "M1 DI"
    assert (sheet["A2"].value, sheet["B2"].value) == ("01", "Links")
    assert (sheet["C2"].value, sheet["D2"].value) == ("02", "Rechts")
    book.close()


def test_document_header_is_copied_with_editable_placeholders(tmp_path: Path) -> None:
    master = tmp_path / "master.xlsx"
    output = tmp_path / "with_header.xlsx"
    make_master(master)
    config = configuration()
    config["output_sheet"] = "{asp}"
    config["document_header"] = {
        "sheet": "Master",
        "range": "A1:D1",
        "row_spacing_after": 1,
        "cell_values": {"A1": "{{DOKUMENTTITEL}}"},
    }
    generate_workbook(master, output, [Record(asp="ASP-1", module="M1", module_type="DI")], config)
    book = load_workbook(output)
    sheet = book["ASP-1"]
    assert sheet["A1"].value == "{{DOKUMENTTITEL}}"
    assert "A1:D1" in {str(item) for item in sheet.merged_cells.ranges}
    assert sheet["A3"].value == "ASP-1 DI"
    book.close()
