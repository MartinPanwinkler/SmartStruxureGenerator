from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


ROWS = [
    ["ASP-301-1", "Modul01", "RTD-DI-16", "In1", "RAV-205.RAUD75.TZU001_MW", "Istwert Ausblastemperatur", "Modicon 7 - M172 - AI1"],
    ["ASP-301-1", "Modul01", "RTD-DI-16", "In2", "RAV-205.RAUD75.TZU002_MW", "Istwert Zulufttemperatur", "Modicon 7 - M172 - AI2"],
    ["ASP-301-1", "Modul01", "RTD-DI-16", "In3", "RAV-205.RAUD75.TZU003_MW", "Istwert Ablufttemperatur", "Modicon 7 - M172 - AI3"],
    ["ASP-303", "Modul02", "DO-FA-12", "Out1", "RLT007.BFS001_EIN_SB", "Schaltbefehl Entrauchung EIN", "ASP0303s_RLT007_EIN_SB"],
]


def create(path: Path = Path("example_input.xlsx")) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Rohdaten"
    headers = ["ASP", "Modul", "Modultyp", "Kanal", "Datenpunkt", "Beschreibung", "Quelle"]
    sheet.append(headers)
    for row in ROWS:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
    widths = [15, 15, 16, 12, 31, 34, 28]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    sheet.freeze_panes = "A2"
    workbook.save(path)


if __name__ == "__main__":
    create()
