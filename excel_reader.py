from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from models import SourceTable


SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm", ".csv"}
SMARTSTRUXURE_HEADERS = [
    "ASP", "Modul", "Modultyp", "Modulnummer", "Kanal",
    "Datenpunkt", "Beschreibung", "Quelle",
]


def list_sheets(path: Path) -> list[str]:
    if path.suffix.lower() == ".csv":
        return []
    workbook = load_workbook(path, read_only=True, data_only=True, keep_vba=path.suffix.lower() == ".xlsm")
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def _csv_encoding(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            path.read_text(encoding=encoding)
            return encoding
        except UnicodeDecodeError:
            continue
    return "latin-1"


def _asp_from_sheet_name(sheet_name: str) -> str:
    upper = sheet_name.upper()
    marker = upper.find("_DP")
    return sheet_name[:marker].strip() if marker >= 0 else sheet_name.strip()


def _looks_like_smartstruxure_export(worksheet: Any) -> bool:
    """Recognize the fixed EBO export layout used by the reference DP sheets."""
    if worksheet.max_column < 7:
        return False
    module_rows = 0
    datapoint_rows = 0
    for b, c, _d, e, f, _g in worksheet.iter_rows(
        min_row=1, max_row=min(worksheet.max_row, 50), min_col=2, max_col=7, values_only=True
    ):
        if b not in (None, "") and c not in (None, "") and e not in (None, "") and f in (None, ""):
            module_rows += 1
        elif c not in (None, "") and f not in (None, ""):
            datapoint_rows += 1
    return module_rows >= 1 and datapoint_rows >= 1


def _load_smartstruxure_export(worksheet: Any) -> SourceTable:
    asp = _asp_from_sheet_name(worksheet.title)
    module = module_type = module_number = ""
    rows: list[dict[str, Any]] = []
    for b, c, d, e, f, g in worksheet.iter_rows(min_col=2, max_col=7, values_only=True):
        is_module = b not in (None, "") and c not in (None, "") and e not in (None, "") and f in (None, "")
        if is_module:
            module_type = str(b).strip()
            module = str(c).strip()
            module_number = str(e).strip()
            continue
        if not module or c in (None, "") or f in (None, ""):
            continue
        rows.append({
            "ASP": asp,
            "Modul": module,
            "Modultyp": module_type,
            "Modulnummer": module_number,
            "Kanal": str(f).strip(),
            "Datenpunkt": str(c).strip(),
            "Beschreibung": "" if d is None else str(d).strip(),
            "Quelle": "" if g is None else str(g).strip(),
        })
    return SourceTable(SMARTSTRUXURE_HEADERS.copy(), rows, list(worksheet.parent.sheetnames), worksheet.title)


def load_source_file(path: Path, sheet_name: str | None = None) -> SourceTable:
    if not path.exists():
        raise FileNotFoundError(f"Rohdatei nicht gefunden: {path}")
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError("Unterstützt werden nur .xlsx, .xlsm und .csv.")
    if suffix == ".csv":
        with path.open("r", encoding=_csv_encoding(path), newline="") as handle:
            sample = handle.read(8192)
            handle.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
            except csv.Error:
                dialect = csv.excel
            reader = csv.DictReader(handle, dialect=dialect)
            headers = [str(value or "").strip() for value in (reader.fieldnames or [])]
            rows = [{str(key or "").strip(): value for key, value in row.items()} for row in reader]
        return SourceTable(headers=headers, rows=rows)

    workbook = load_workbook(path, read_only=True, data_only=True, keep_vba=suffix == ".xlsm")
    try:
        sheets = list(workbook.sheetnames)
        selected = sheet_name or sheets[0]
        if selected not in sheets:
            raise ValueError(f"Tabellenblatt nicht gefunden: {selected}")
        worksheet = workbook[selected]
        if _looks_like_smartstruxure_export(worksheet):
            return _load_smartstruxure_export(worksheet)
        values = worksheet.iter_rows(values_only=True)
        first = next(values, None)
        if first is None:
            return SourceTable([], [], sheets, selected)
        headers = [str(value).strip() if value is not None else f"Spalte_{index + 1}" for index, value in enumerate(first)]
        rows: list[dict[str, Any]] = []
        for values_row in values:
            row = {headers[index]: value for index, value in enumerate(values_row) if index < len(headers)}
            if any(value not in (None, "") for value in row.values()):
                rows.append(row)
        return SourceTable(headers, rows, sheets, selected)
    finally:
        workbook.close()
