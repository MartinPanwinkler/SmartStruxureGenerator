from __future__ import annotations

import logging
from copy import copy
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import column_index_from_string, get_column_letter

from config import BOX_ROW_SPACING, GROUP_BY
from mapping import group_records, split_records
from models import BoxTemplate, GenerationResult, ProgressCallback, Record
from template_analyzer import PLACEHOLDER, load_templates, select_template
from template_analyzer import box_from_config


LOGGER = logging.getLogger(__name__)


def _copy_cell(source: Any, target: Any) -> None:
    if not isinstance(source, MergedCell):
        target.value = source.value
    if source.has_style:
        # Style IDs are workbook-local; copy components so openpyxl registers
        # each style in the destination workbook.
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
    target.number_format = source.number_format
    target.protection = copy(source.protection)
    target.alignment = copy(source.alignment)
    if not isinstance(source, MergedCell) and source.hyperlink:
        target._hyperlink = copy(source.hyperlink)
    if not isinstance(source, MergedCell) and source.comment:
        target.comment = copy(source.comment)


def copy_box(source_sheet: Any, target_sheet: Any, template: BoxTemplate, target_row: int, target_col: int = 1) -> tuple[int, int]:
    row_offset = target_row - template.min_row
    col_offset = target_col - template.min_col
    for row in range(template.min_row, template.max_row + 1):
        source_dimension = source_sheet.row_dimensions[row]
        target_dimension = target_sheet.row_dimensions[row + row_offset]
        target_dimension.height = source_dimension.height
        target_dimension.hidden = source_dimension.hidden
        for col in range(template.min_col, template.max_col + 1):
            source_cell = source_sheet.cell(row, col)
            _copy_cell(source_cell, target_sheet.cell(row + row_offset, col + col_offset))
    for col in range(template.min_col, template.max_col + 1):
        source_letter = get_column_letter(col)
        target_letter = get_column_letter(col + col_offset)
        source_dimension = source_sheet.column_dimensions[source_letter]
        target_dimension = target_sheet.column_dimensions[target_letter]
        target_dimension.width = source_dimension.width
        target_dimension.hidden = source_dimension.hidden
    for merged in source_sheet.merged_cells.ranges:
        if (merged.min_row >= template.min_row and merged.max_row <= template.max_row
                and merged.min_col >= template.min_col and merged.max_col <= template.max_col):
            target_sheet.merge_cells(
                start_row=merged.min_row + row_offset,
                end_row=merged.max_row + row_offset,
                start_column=merged.min_col + col_offset,
                end_column=merged.max_col + col_offset,
            )
    return template.max_row + row_offset, template.max_col + col_offset


def _copy_sheet_settings(source: Any, target: Any) -> None:
    target.sheet_format = copy(source.sheet_format)
    target.sheet_properties = copy(source.sheet_properties)
    target.page_margins = copy(source.page_margins)
    target.page_setup = copy(source.page_setup)
    target.print_options = copy(source.print_options)
    target.freeze_panes = source.freeze_panes
    target.sheet_view.showGridLines = source.sheet_view.showGridLines


def _copy_document_header(master: Any, target: Any, config: dict[str, Any]) -> int:
    header_config = config.get("document_header")
    if not isinstance(header_config, dict) or not header_config.get("enabled", True):
        return 1
    parse_config = dict(header_config)
    parse_config.setdefault("max_channels", 1)
    header = box_from_config("document_header", parse_config)
    if header.sheet not in master.sheetnames:
        raise ValueError(f"Master-Tabellenblatt für Dokumentkopf nicht gefunden: {header.sheet}")
    source = master[header.sheet]
    _copy_sheet_settings(source, target)
    copy_box(source, target, header, target_row=1, target_col=header.min_col)
    for coordinate, value in dict(header_config.get("cell_values", {})).items():
        row, col = _relative_coordinate(coordinate, header, 1, header.min_col)
        cell = target.cell(row, col)
        if isinstance(cell, MergedCell):
            raise ValueError(f"Dokumentkopf: Zielzelle {coordinate} ist keine Merge-Ankerzelle.")
        cell.value = value
    return header.height + int(header_config.get("row_spacing_after", 1)) + 1


def _relative_coordinate(coordinate: str, template: BoxTemplate, target_row: int, target_col: int) -> tuple[int, int]:
    from openpyxl.utils.cell import coordinate_from_string
    column, row = coordinate_from_string(coordinate)
    return target_row + row - template.min_row, target_col + column_index_from_string(column) - template.min_col


def _replace_placeholders(value: Any, values: dict[str, Any]) -> Any:
    if not isinstance(value, str):
        return value
    full = PLACEHOLDER.fullmatch(value.strip())
    if full:
        return values.get(full.group(1).lower(), "")
    return PLACEHOLDER.sub(lambda match: str(values.get(match.group(1).lower(), "")), value)


def populate_box(sheet: Any, template: BoxTemplate, target_row: int, records: list[Record], target_col: int = 1) -> None:
    group_values = {key: records[0].value(key) if records else "" for key in Record.__dataclass_fields__ if key != "extra"}
    module_parts = [group_values.get("module", ""), group_values.get("module_type", "")]
    group_values["module_header"] = " ".join(str(value) for value in module_parts if value).strip()
    for field, format_string in template.header_formats.items():
        try:
            group_values[field] = format_string.format_map(group_values)
        except KeyError as exc:
            raise ValueError(f"Template '{template.name}': unbekanntes Feld in header_formats: {exc}") from exc
    for field, configured_coordinates in template.header_cells.items():
        coordinates = [configured_coordinates] if isinstance(configured_coordinates, str) else configured_coordinates
        for coordinate in coordinates:
            row, col = _relative_coordinate(coordinate, template, target_row, target_col)
            sheet.cell(row, col).value = group_values.get(field, "")
    for row in range(target_row, target_row + template.height):
        for col in range(target_col, target_col + template.width):
            cell = sheet.cell(row, col)
            if not isinstance(cell, MergedCell):
                cell.value = _replace_placeholders(cell.value, group_values)
    if template.data_slots:
        if len(template.data_slots) < template.max_channels:
            raise ValueError(f"Template '{template.name}' definiert weniger data_slots als max_channels.")
        for index in range(template.max_channels):
            values = records[index] if index < len(records) else Record()
            for field, coordinate in template.data_slots[index].items():
                if field in template.preserve_template_fields:
                    continue
                row, col = _relative_coordinate(coordinate, template, target_row, target_col)
                cell = sheet.cell(row, col)
                if isinstance(cell, MergedCell):
                    raise ValueError(f"Template '{template.name}': Datenziel {coordinate} ist verbunden.")
                cell.value = values.value(field)
    else:
        data_offset = template.data_start_row - template.min_row
        for index in range(template.max_channels):
            values = records[index] if index < len(records) else Record()
            destination_row = target_row + data_offset + index
            if destination_row >= target_row + template.height:
                raise ValueError(f"Template '{template.name}' hat nicht genug Datenzeilen für max_channels.")
            for field, column in template.columns.items():
                source_col = column_index_from_string(column)
                destination_col = target_col + source_col - template.min_col
                cell = sheet.cell(destination_row, destination_col)
                if isinstance(cell, MergedCell):
                    raise ValueError(f"Template '{template.name}': Datenziel {column}{template.data_start_row + index} ist verbunden.")
                cell.value = values.value(field)


def generate_workbook(
    master_path: Path,
    output_path: Path,
    records: list[Record],
    config: dict[str, Any],
    progress: ProgressCallback | None = None,
) -> GenerationResult:
    master_path = master_path.resolve()
    output_path = output_path.resolve()
    if master_path == output_path:
        raise ValueError("Die Master-Vorlage darf nicht überschrieben werden.")
    if output_path.exists():
        raise FileExistsError(f"Die Ausgabedatei existiert bereits: {output_path}")
    templates = load_templates(config)
    if not templates:
        raise ValueError("In template_config.json ist kein verwendbares Template definiert.")
    mapping = dict(config.get("template_mapping", {}))
    default = str(config.get("default_template", next(iter(templates))))
    group_by = list(config.get("group_by", GROUP_BY))
    spacing = int(config.get("box_row_spacing", BOX_ROW_SPACING))
    start_col = int(config.get("output_start_col", 1))
    groups = group_records(records, group_by)
    master = load_workbook(master_path, data_only=False)
    output = Workbook()
    target = output.active
    first_values = {key: records[0].value(key) if records else "" for key in Record.__dataclass_fields__ if key != "extra"}
    sheet_pattern = str(config.get("output_sheet", "Beschriftungen"))
    try:
        sheet_name = sheet_pattern.format_map(first_values)
    except KeyError:
        sheet_name = sheet_pattern
    target.title = (sheet_name or "Beschriftungen")[:31]
    boxes = 0
    current_row = 1
    try:
        current_row = _copy_document_header(master, target, config)
        total_groups = max(len(groups), 1)
        for group_index, (_, group) in enumerate(groups):
            module_type = group[0].module_type if group else ""
            template = select_template(module_type, templates, mapping, default)
            if template.sheet not in master.sheetnames:
                raise ValueError(f"Master-Tabellenblatt nicht gefunden: {template.sheet}")
            for chunk in split_records(group, template.max_channels):
                if boxes == 0 and current_row == 1:
                    _copy_sheet_settings(master[template.sheet], target)
                copy_box(master[template.sheet], target, template, current_row, start_col)
                populate_box(target, template, current_row, chunk, start_col)
                current_row += template.height + spacing
                boxes += 1
            if progress:
                progress(round((group_index + 1) / total_groups * 100), f"Gruppe {group_index + 1} von {len(groups)}")
        if boxes:
            print_start_col = 1 if isinstance(config.get("document_header"), dict) else start_col
            target.print_area = (
                f"{get_column_letter(print_start_col)}1:"
                f"{get_column_letter(target.max_column)}{target.max_row}"
            )
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output.save(output_path)
    finally:
        master.close()
        output.close()
    warnings = []
    missing_descriptions = sum(not record.description for record in records)
    if missing_descriptions:
        warnings.append(f"Bei {missing_descriptions} Datenpunkten fehlt eine Beschreibung.")
    LOGGER.info("Ausgabe erstellt: %s (%d Kästchen, %d Datensätze)", output_path, boxes, len(records))
    return GenerationResult(output_path, boxes, len(records), warnings)
