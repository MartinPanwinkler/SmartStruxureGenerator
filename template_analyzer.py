from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter, range_boundaries

from models import BoxTemplate


PLACEHOLDER = re.compile(r"\{\{\s*([A-Za-z0-9_]+)\s*\}\}")


def _used(cell: Any) -> bool:
    if isinstance(cell, MergedCell):
        return False
    border = cell.border
    bordered = any(getattr(border, side).style for side in ("left", "right", "top", "bottom"))
    return cell.value not in (None, "") or bordered or cell.fill.fill_type is not None


def detect_box_ranges(worksheet: Any) -> list[str]:
    """Detect visual row blocks separated by completely unused rows.

    This deliberately returns candidates, not semantic template names. A human can
    confirm or correct them in template_config.json.
    """
    active_rows: list[int] = []
    extents: dict[int, tuple[int, int]] = {}
    for row in range(1, worksheet.max_row + 1):
        columns = [cell.column for cell in worksheet[row] if _used(cell)]
        for merged in worksheet.merged_cells.ranges:
            if merged.min_row <= row <= merged.max_row:
                columns.extend((merged.min_col, merged.max_col))
        if columns:
            active_rows.append(row)
            extents[row] = (min(columns), max(columns))
    if not active_rows:
        return []
    bands: list[list[int]] = [[active_rows[0]]]
    for row in active_rows[1:]:
        if row > bands[-1][-1] + 1:
            bands.append([row])
        else:
            bands[-1].append(row)
    ranges: list[str] = []
    for band in bands:
        min_col = min(extents[row][0] for row in band)
        max_col = max(extents[row][1] for row in band)
        ranges.append(f"{get_column_letter(min_col)}{band[0]}:{get_column_letter(max_col)}{band[-1]}")
    return ranges


def box_from_config(name: str, data: dict[str, Any]) -> BoxTemplate:
    required = {"sheet", "range", "max_channels"}
    missing = sorted(required - data.keys())
    if missing:
        raise ValueError(f"Template '{name}': Pflichtangaben fehlen: {', '.join(missing)}")
    min_col, min_row, max_col, max_row = range_boundaries(data["range"])
    start = int(data.get("data_start_row", min_row))
    if start < min_row or start > max_row:
        raise ValueError(f"Template '{name}': data_start_row liegt außerhalb des Bereichs.")
    return BoxTemplate(
        name=name,
        sheet=str(data["sheet"]),
        min_row=min_row,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
        header_cells=dict(data.get("header_cells", {})),
        data_start_row=start,
        max_channels=int(data["max_channels"]),
        columns=dict(data.get("columns", {})),
        data_slots=[dict(slot) for slot in data.get("data_slots", [])],
        header_formats=dict(data.get("header_formats", {})),
        preserve_template_fields=set(data.get("preserve_template_fields", [])),
    )


def load_templates(config: dict[str, Any]) -> dict[str, BoxTemplate]:
    raw = config.get("templates", config)
    return {name: box_from_config(name, value) for name, value in raw.items() if isinstance(value, dict) and "range" in value}


def select_template(module_type: str, templates: dict[str, BoxTemplate], mapping: dict[str, str], default: str) -> BoxTemplate:
    lookup = module_type.strip().casefold()
    mapped = next((target for key, target in mapping.items() if key.casefold() == lookup), None)
    name = mapped or next((key for key in templates if key.casefold() == lookup), default)
    if name not in templates:
        raise ValueError(f"Template '{name}' ist nicht konfiguriert.")
    return templates[name]


def analyze_template(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=False)
    try:
        report: dict[str, Any] = {"file": str(path), "sheets": {}}
        for worksheet in workbook.worksheets:
            values = []
            placeholders = []
            for row in worksheet.iter_rows():
                for cell in row:
                    if not isinstance(cell, MergedCell) and cell.value not in (None, ""):
                        values.append({"cell": cell.coordinate, "value": str(cell.value)})
                        if isinstance(cell.value, str):
                            placeholders.extend({"cell": cell.coordinate, "field": match.group(1).lower()} for match in PLACEHOLDER.finditer(cell.value))
            report["sheets"][worksheet.title] = {
                "dimensions": worksheet.calculate_dimension(),
                "candidate_boxes": detect_box_ranges(worksheet),
                "merged_cells": [str(item) for item in worksheet.merged_cells.ranges],
                "row_heights": {str(index): dim.height for index, dim in worksheet.row_dimensions.items() if dim.height is not None},
                "column_widths": {key: dim.width for key, dim in worksheet.column_dimensions.items() if dim.width is not None},
                "cells_with_values": values,
                "placeholders": placeholders,
            }
        return report
    finally:
        workbook.close()


def write_analysis(path: Path, output: Path) -> None:
    output.write_text(json.dumps(analyze_template(path), ensure_ascii=False, indent=2), encoding="utf-8")
